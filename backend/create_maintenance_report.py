"""Build the Maintenance Photo Report DOCX (letterhead cover + checklist + photo grid pages)."""
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from create_maintenance_checklist import add_checklist_section

PHOTOS_PER_PAGE = 4  # laid out as a 2x2 grid, matching the sample report's density


def _format_session_dates(dates: list[str]) -> tuple[str, str]:
    """
    dates: 'YYYY-MM-DD' strings from the selected photos (may span one month).
    Returns (month_label e.g. "May 2026", detail_label e.g. "26 and 28 May 2026").
    Falls back to a plain min-max range if dates span more than one month.
    """
    parsed = sorted({datetime.strptime(d, "%Y-%m-%d") for d in dates})
    if not parsed:
        return "", ""

    months = {(p.year, p.month) for p in parsed}
    if len(months) > 1:
        rng = f"{parsed[0].strftime('%d %b %Y')} to {parsed[-1].strftime('%d %b %Y')}"
        return rng, rng

    month_label = parsed[-1].strftime("%B %Y")
    days = [str(p.day) for p in parsed]
    if len(days) == 1:
        day_str = days[0]
    elif len(days) == 2:
        day_str = f"{days[0]} and {days[1]}"
    else:
        day_str = ", ".join(days[:-1]) + f" and {days[-1]}"
    return month_label, f"{day_str} {month_label}"


def _set_landscape(doc):
    for section in doc.sections:
        section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = 1  # WD_ORIENT.LANDSCAPE
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)


def _prevent_row_split(row):
    """Stop a table row (image + caption) from splitting across a page boundary --
    without this, a photo can render on one page while its caption spills onto
    the next."""
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def _add_gws_header(doc, logo_bytes):
    table = doc.add_table(rows=1, cols=2)
    logo_cell, text_cell = table.rows[0].cells
    if logo_bytes:
        logo_cell.paragraphs[0].add_run().add_picture(io.BytesIO(logo_bytes), height=Cm(1.6))

    p = text_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("GWS Living Art PTE LTD")
    run.bold = True
    run.font.size = Pt(12)

    p2 = text_cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("Reg no.: 201709254M")
    r2.font.size = Pt(8)

    doc.add_paragraph()


def build_report_docx(
    project: dict,
    maintenance_date: str,
    photos: list[dict],
    logo_bytes: bytes = None,
) -> bytes:
    """
    photos: list of {"bytes": <image bytes>, "date": "YYYY-MM-DD"} in the order
    they should appear in the report. Cover/header date labels are derived
    from these photos' own dates, not a separately-passed range.
    """
    month_label, detail_label = _format_session_dates([p["date"] for p in photos if p.get("date")])

    doc = Document()
    _set_landscape(doc)

    # ── Cover page ───────────────────────────────────────────────────────────
    cover = doc.add_table(rows=1, cols=2)
    client_cell, logo_cell = cover.rows[0].cells

    client_lines = [line for line in [project.get("client"), project.get("address")] if line]
    if client_lines:
        cp = client_cell.paragraphs[0]
        run = cp.add_run(client_lines[0])
        run.bold = True
        run.font.size = Pt(13)
        for line in client_lines[1:]:
            lp = client_cell.add_paragraph()
            lp.add_run(line).font.size = Pt(10)

    if logo_bytes:
        lp = logo_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lp.add_run().add_picture(io.BytesIO(logo_bytes), height=Cm(2.2))

    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"GREENROOF MAINTENANCE PHOTO REPORT FOR\n{project.get('name', '').upper()}")
    run.bold = True
    run.font.size = Pt(20)

    for _ in range(3):
        doc.add_paragraph()

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"Maintenance sessions for {month_label}")
    r.bold = True
    r.font.size = Pt(12)

    doc.add_page_break()

    # ── Checklist page (included in the report, not a separate document) ────
    _add_gws_header(doc, logo_bytes)
    add_checklist_section(doc, project, maintenance_date)

    doc.add_page_break()

    # ── Session header page ─────────────────────────────────────────────────
    _add_gws_header(doc, logo_bytes)
    for _ in range(3):
        doc.add_paragraph()

    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = header.add_run(f"GREEN ROOF MAINTENANCE – {detail_label}")
    hr.bold = True
    hr.font.size = Pt(14)

    doc.add_page_break()

    # ── Photo grid pages ─────────────────────────────────────────────────────
    for page_start in range(0, len(photos), PHOTOS_PER_PAGE):
        _add_gws_header(doc, logo_bytes)
        page_photos = photos[page_start:page_start + PHOTOS_PER_PAGE]

        rows = (len(page_photos) + 1) // 2
        grid = doc.add_table(rows=rows, cols=2)
        for i, photo in enumerate(page_photos):
            row = grid.rows[i // 2]
            _prevent_row_split(row)
            cell = row.cells[i % 2]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(photo["bytes"]), width=Cm(11))
            cap = cell.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            date_str = photo.get("date", "")
            try:
                date_label = datetime.strptime(date_str, "%Y-%m-%d").strftime("%d %b %Y")
            except ValueError:
                date_label = date_str
            cap.add_run(date_label).font.size = Pt(8)

        if page_start + PHOTOS_PER_PAGE < len(photos):
            doc.add_page_break()

    # ── End page ─────────────────────────────────────────────────────────────
    doc.add_page_break()
    _add_gws_header(doc, logo_bytes)
    for _ in range(4):
        doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = end.add_run("Thank You")
    er.bold = True
    er.font.size = Pt(24)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
